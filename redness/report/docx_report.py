"""findings.json / Report  ->  minimal, classic network-pentest .docx.

Person B's domain. No corporate template - plain, table-driven, readable:
  - Title + engagement meta
  - Executive summary (severity counts table)
  - Findings, most-severe first, each with impact / remediation / evidence / affected
  - Appendix: enumeration (informational)

Depends on python-docx.
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..findings import Category, Report, Severity

_SEV_COLOR = {
    Severity.CRITICAL: RGBColor(0x8B, 0x00, 0x00),
    Severity.HIGH: RGBColor(0xC0, 0x39, 0x2B),
    Severity.MEDIUM: RGBColor(0xB9, 0x77, 0x0E),
    Severity.LOW: RGBColor(0x1F, 0x6F, 0x3C),
    Severity.INFO: RGBColor(0x55, 0x55, 0x55),
}
_SEV_LABEL = {s: s.value.upper() for s in Severity}


def _shade(cell, hex_fill: str) -> None:
    """Apply a solid background fill to a table cell."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def build(report: Report, out_path: str) -> str:
    doc = Document()

    # ---- Title -------------------------------------------------------------
    title = doc.add_heading("Active Directory Security Assessment", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Engagement: {report.engagement}\n").bold = True
    meta.add_run(f"Target: {report.target or 'N/A'}")

    # ---- Executive summary -------------------------------------------------
    doc.add_heading("Executive Summary", level=1)
    counts = {s: len(v) for s, v in report.by_severity().items()}
    total = sum(counts.values())
    doc.add_paragraph(
        f"This assessment identified {total} finding(s) across the in-scope "
        f"Active Directory environment. The table below summarizes findings by severity."
    )

    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text, hdr[1].text = "Severity", "Count"
    for s in [Severity.CRITICAL, Severity.HIGH, Severity.MEDIUM, Severity.LOW, Severity.INFO]:
        row = t.add_row().cells
        run = row[0].paragraphs[0].add_run(_SEV_LABEL[s])
        run.bold = True
        run.font.color.rgb = _SEV_COLOR[s]
        row[1].text = str(counts[s])

    # ---- Findings ----------------------------------------------------------
    doc.add_heading("Findings", level=1)
    ordered = [f for f in report.sorted_findings() if f.category != Category.ENUM]
    if not ordered:
        doc.add_paragraph("No security findings recorded.")
    for i, f in enumerate(ordered, 1):
        h = doc.add_heading(level=2)
        run = h.add_run(f"{i}. {f.title}")
        run.font.color.rgb = _SEV_COLOR[f.severity]

        # metadata mini-table
        mt = doc.add_table(rows=0, cols=2)
        mt.style = "Light List Accent 1"
        for label, value in [
            ("Severity", _SEV_LABEL[f.severity]),
            ("Category", f.category.value.upper()),
            ("Access", "Credentialed" if f.credentialed else "Unauthenticated"),
            ("Affected", ", ".join(f.affected) if f.affected else "-"),
        ]:
            r = mt.add_row().cells
            r[0].paragraphs[0].add_run(label).bold = True
            r[1].text = value

        if f.description:
            doc.add_paragraph().add_run("Description").bold = True
            doc.add_paragraph(f.description)
        if f.impact:
            doc.add_paragraph().add_run("Impact").bold = True
            doc.add_paragraph(f.impact)
        if f.remediation:
            doc.add_paragraph().add_run("Remediation").bold = True
            doc.add_paragraph(f.remediation)
        if f.evidence:
            doc.add_paragraph().add_run("Evidence").bold = True
            for ev in f.evidence:
                p = doc.add_paragraph(style="List Bullet")
                p.add_run(f"[{ev.source}] ").bold = True
                p.add_run(ev.detail)
                if ev.raw:
                    code = doc.add_paragraph(ev.raw)
                    code.style = doc.styles["No Spacing"]
                    for cr in code.runs:
                        cr.font.name = "Consolas"
                        cr.font.size = Pt(8)
        if f.references:
            doc.add_paragraph().add_run("References").bold = True
            for ref in f.references:
                doc.add_paragraph(ref, style="List Bullet")

    # ---- Appendix: enumeration --------------------------------------------
    enum = [f for f in report.sorted_findings() if f.category == Category.ENUM]
    if enum:
        doc.add_heading("Appendix A - Enumeration", level=1)
        for f in enum:
            doc.add_heading(f.title, level=2)
            if f.affected:
                for a in f.affected:
                    doc.add_paragraph(a, style="List Bullet")
            if f.description:
                doc.add_paragraph(f.description)

    doc.save(out_path)
    return out_path
