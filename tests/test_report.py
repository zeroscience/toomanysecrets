"""Report generation tests: Attack Paths section + MITRE ATT&CK mapping.

Needs python-docx (a base dependency). Run: python tests/test_report.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from redness.report import docx_report  # noqa: E402
from redness.report.mappings import compliance_for, mitre_for  # noqa: E402
from redness.findings import Category, Finding, Report, Severity  # noqa: E402


def _report():
    r = Report(engagement="Lab", target="acme.local")
    r.findings = [
        Finding(id="krb-asrep-roastable", title="AS-REP roastable", severity=Severity.HIGH,
                category=Category.KERBEROS, affected=["jdoe", "svc-sql"]),
        Finding(id="credential-cracked", title="Password recovered for svc-sql",
                severity=Severity.CRITICAL, category=Category.KERBEROS, affected=["svc-sql"],
                meta={"account": "svc-sql", "password": "Summer2024!"}),
        Finding(id="adcs-esc1", title="AD CS ESC1", severity=Severity.CRITICAL,
                category=Category.ADCS, affected=["UserAuth"], meta={"esc": 1}),
        Finding(id="smb-signing-not-required", title="SMB signing not required",
                severity=Severity.MEDIUM, category=Category.SMB, affected=["10.0.0.1"]),
    ]
    return r


def _docx_text(path):
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_mitre_mapping():
    assert mitre_for("krb-asrep-roastable")[0] == "T1558.004"
    assert mitre_for("krb-kerberoastable")[0] == "T1558.003"
    assert mitre_for("adcs-esc7")[0] == "T1649"      # any ESC -> T1649
    assert mitre_for("tls-legacy-protocol") is None  # unmapped is fine


def test_report_has_attack_paths_and_mitre():
    with tempfile.TemporaryDirectory() as td:
        out = os.path.join(td, "r.docx")
        docx_report.build(_report(), out)
        text = _docx_text(out)
    assert "Attack Paths" in text
    assert "Compromised (owned) accounts: svc-sql" in text
    assert "Crackable credentials (roastable): jdoe" in text          # svc-sql owned, excluded
    assert "Direct path to Domain Admin: UserAuth via AD CS ESC1" in text
    assert "T1558.004" in text and "T1649" in text                    # MITRE surfaced


def test_compliance_mapping():
    c = compliance_for("krb-kerberoastable")
    assert c and set(c) == {"CIS", "ISO", "NIS2"}
    assert compliance_for("adcs-esc7")["ISO"].startswith("A.")      # ADCS prefix
    assert compliance_for("tls-legacy-protocol")["CIS"].startswith("3.10")  # TLS prefix
    assert compliance_for("enum-domain-users") is None              # unmapped is fine


def test_report_has_compliance_appendix():
    with tempfile.TemporaryDirectory() as td:
        out = os.path.join(td, "r.docx")
        docx_report.build(_report(), out)
        text = _docx_text(out)
    assert "Appendix B - Compliance Mapping" in text
    assert "CIS v8" in text and "ISO 27001:2022" in text and "NIS2" in text
    assert "A.8.5 Secure authentication" in text                    # from a mapped finding


def test_no_attack_paths_section_without_paths():
    r = Report(engagement="Lab", target="x")
    r.findings = [Finding(id="enum-domain-users", title="users", severity=Severity.INFO,
                          category=Category.ENUM, affected=["a", "b"])]
    with tempfile.TemporaryDirectory() as td:
        out = os.path.join(td, "r.docx")
        docx_report.build(r, out)
        text = _docx_text(out)
    assert "Attack Paths" not in text


if __name__ == "__main__":
    passed = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print(f"  ok  {name}")
            passed += 1
    print(f"\n{passed} tests passed")
